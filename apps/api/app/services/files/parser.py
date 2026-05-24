from pathlib import Path

import pytesseract
from bs4 import BeautifulSoup
from docx import Document
from PIL import Image
from pypdf import PdfReader


class FileParser:
    def parse(self, path: Path, content_type: str | None) -> tuple[str | None, dict]:
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md", ".markdown"}:
            text = path.read_text(encoding="utf-8", errors="ignore")
            return text, {"parser": "plain-text", "characters": len(text)}
        if suffix == ".docx":
            document = Document(path)
            text = "\n".join(
                paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
            )
            return text, {"parser": "python-docx", "paragraphs": len(document.paragraphs)}
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            pages: list[str] = []
            for page in reader.pages:
                pages.append(page.extract_text() or "")
            text = "\n".join(part for part in pages if part.strip())
            return text or None, {"parser": "pypdf", "pages": len(reader.pages)}
        if content_type and content_type.startswith("image/"):
            image = Image.open(path)
            text = pytesseract.image_to_string(image)
            return text or None, {"parser": "pytesseract", "mode": "ocr"}
        if suffix in {".html", ".htm"}:
            html = path.read_text(encoding="utf-8", errors="ignore")
            soup = BeautifulSoup(html, "html.parser")
            return soup.get_text("\n", strip=True), {"parser": "beautifulsoup", "tags": len(soup.find_all())}
        return None, {"parser": "unsupported"}
